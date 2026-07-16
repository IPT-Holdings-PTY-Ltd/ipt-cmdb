"""Governance report catalogue and branded export renderers."""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from typing import Any


REPORT_CATALOG = [
    {"id": "asset-register", "title": "Asset register", "description": "Canonical CI inventory, source, lifecycle, health and ownership.", "rootOnly": False},
    {"id": "lifecycle-attention", "title": "Lifecycle and renewals", "description": "Renewals, warranty expiry, end-of-life dates and review dates requiring attention.", "rootOnly": False},
    {"id": "ownership-gaps", "title": "Ownership gaps", "description": "Configuration items missing business, service or technical accountability.", "rootOnly": False},
    {"id": "business-services", "title": "Business-system register", "description": "Business services, accountable owners, recovery objectives and operating state.", "rootOnly": False},
    {"id": "change-register", "title": "Change register", "description": "Planned and generated changes with risk, schedule, owner and impact totals.", "rootOnly": False},
    {"id": "audit-activity", "title": "Audit activity", "description": "Attributable governance events with outcome, source and correlation reference.", "rootOnly": False},
    {"id": "access-review", "title": "Effective access review", "description": "Users, effective roles and customer scope for periodic access certification.", "rootOnly": True},
    {"id": "integration-health", "title": "Integration health", "description": "Connection readiness and recent synchronisation outcomes.", "rootOnly": True},
]


def report_catalog(root_scope: bool) -> list[dict]:
    return [{key: value for key, value in item.items() if key != "rootOnly"} for item in REPORT_CATALOG if root_scope or not item["rootOnly"]]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (list, tuple, set)):
        return ", ".join(_text(item) for item in value)
    if isinstance(value, dict):
        return "; ".join(f"{key}: {_text(child)}" for key, child in value.items())
    return str(value)


def _columns(*items: tuple[str, str]) -> list[dict]:
    return [{"key": key, "label": label} for key, label in items]


def build_report(
    report_id: str,
    *,
    companies: list[dict],
    assets: list[dict],
    relationships: list[dict],
    changes: list[dict],
    users: list[dict],
    integrations: list[dict],
    sync_runs: list[dict],
    audit_events: list[dict],
    company_id: str | None,
) -> dict:
    definition = next((item for item in REPORT_CATALOG if item["id"] == report_id), None)
    if not definition:
        raise ValueError("Unknown report")
    company_names = {item["id"]: item["name"] for item in companies}
    scoped_assets = [item for item in assets if not company_id or item["companyId"] == company_id]
    scoped_changes = [item for item in changes if not company_id or item["companyId"] == company_id]
    rows: list[dict] = []

    if report_id == "asset-register":
        columns = _columns(("customer", "Customer"), ("name", "Configuration item"), ("type", "Type"), ("layer", "Layer"), ("status", "Status"), ("lifecycle", "Lifecycle"), ("criticality", "Criticality"), ("owner", "Primary owner"), ("source", "Source"), ("lastSeen", "Last seen"))
        for item in scoped_assets:
            meta = item.get("metadata") or {}
            rows.append({"customer": company_names.get(item["companyId"], item["companyId"]), "name": item["name"], "type": item["type"], "layer": meta.get("displayLayer", ""), "status": meta.get("operationalStatus", item.get("status", "")), "lifecycle": meta.get("lifecycle", ""), "criticality": meta.get("criticality", ""), "owner": meta.get("businessOwner") or meta.get("serviceOwner") or meta.get("technicalOwner") or meta.get("custodian") or "Unassigned", "source": item.get("source", ""), "lastSeen": item.get("lastSeen", "")})
    elif report_id == "lifecycle-attention":
        columns = _columns(("customer", "Customer"), ("name", "Configuration item"), ("type", "Type"), ("lifecycle", "Lifecycle"), ("renewal", "Renewal"), ("warranty", "Warranty end"), ("endOfLife", "End of life"), ("review", "Review date"), ("owner", "Owner"))
        for item in scoped_assets:
            meta = item.get("metadata") or {}
            if not any(meta.get(key) for key in ("renewalDate", "warrantyEnd", "endOfLifeDate", "reviewDate")) and meta.get("lifecycle") not in {"maintenance", "retired", "disposed"}:
                continue
            rows.append({"customer": company_names.get(item["companyId"], item["companyId"]), "name": item["name"], "type": item["type"], "lifecycle": meta.get("lifecycle", ""), "renewal": meta.get("renewalDate", ""), "warranty": meta.get("warrantyEnd", ""), "endOfLife": meta.get("endOfLifeDate", ""), "review": meta.get("reviewDate", ""), "owner": meta.get("serviceOwner") or meta.get("technicalOwner") or "Unassigned"})
    elif report_id == "ownership-gaps":
        columns = _columns(("customer", "Customer"), ("name", "Configuration item"), ("type", "Type"), ("criticality", "Criticality"), ("businessOwner", "Business owner"), ("serviceOwner", "Service owner"), ("technicalOwner", "Technical owner"), ("missing", "Missing accountability"))
        for item in scoped_assets:
            meta = item.get("metadata") or {}
            required = ["businessOwner"] if item.get("type") == "Business system" else []
            required += ["serviceOwner", "technicalOwner"]
            missing = [key.replace("Owner", " owner") for key in required if not meta.get(key)]
            if missing:
                rows.append({"customer": company_names.get(item["companyId"], item["companyId"]), "name": item["name"], "type": item["type"], "criticality": meta.get("criticality", ""), "businessOwner": meta.get("businessOwner", ""), "serviceOwner": meta.get("serviceOwner", ""), "technicalOwner": meta.get("technicalOwner", ""), "missing": missing})
    elif report_id == "business-services":
        columns = _columns(("customer", "Customer"), ("name", "Business system"), ("criticality", "Criticality"), ("health", "Operational status"), ("businessOwner", "Business owner"), ("serviceOwner", "Service owner"), ("rto", "RTO hours"), ("rpo", "RPO hours"), ("supportingCIs", "Supporting CIs"))
        relationship_counts: dict[str, int] = {}
        for relationship in relationships:
            relationship_counts[relationship["fromId"]] = relationship_counts.get(relationship["fromId"], 0) + 1
        for item in scoped_assets:
            if item.get("type") != "Business system":
                continue
            meta = item.get("metadata") or {}
            rows.append({"customer": company_names.get(item["companyId"], item["companyId"]), "name": item["name"], "criticality": meta.get("criticality", ""), "health": meta.get("operationalStatus", ""), "businessOwner": meta.get("businessOwner", ""), "serviceOwner": meta.get("serviceOwner", ""), "rto": meta.get("rtoHours", ""), "rpo": meta.get("rpoHours", ""), "supportingCIs": relationship_counts.get(item["id"], 0)})
    elif report_id == "change-register":
        columns = _columns(("customer", "Customer"), ("number", "Reference"), ("title", "Change"), ("status", "Status"), ("risk", "Risk"), ("start", "Planned start"), ("technician", "Technician"), ("approver", "Approver"), ("impact", "Impacted CIs"))
        for item in scoped_changes:
            rows.append({"customer": company_names.get(item["companyId"], item["companyId"]), "number": item.get("number", ""), "title": item.get("title", ""), "status": item.get("status", ""), "risk": item.get("riskLevel", ""), "start": item.get("plannedStart", ""), "technician": item.get("assignedTechnician", ""), "approver": item.get("approver", ""), "impact": len(item.get("impactSnapshot") or [])})
    elif report_id == "audit-activity":
        columns = _columns(("time", "Time UTC"), ("customer", "Customer"), ("actor", "Actor"), ("category", "Category"), ("action", "Action"), ("entity", "Entity"), ("outcome", "Outcome"), ("source", "Source"), ("correlation", "Correlation ID"))
        for item in audit_events:
            rows.append({"time": item.get("createdAt", ""), "customer": company_names.get(item.get("companyId"), item.get("companyId") or "MSP"), "actor": item.get("actorLabel", "System"), "category": item.get("category", ""), "action": item.get("action", ""), "entity": item.get("entityName") or item.get("entityType", ""), "outcome": item.get("outcome", ""), "source": item.get("sourceSystem", ""), "correlation": item.get("correlationId", "")})
    elif report_id == "access-review":
        columns = _columns(("email", "User"), ("role", "Effective role"), ("accountType", "Account type"), ("customers", "Customer scope"), ("groups", "Access groups"))
        for item in users:
            rows.append({"email": item.get("email", ""), "role": item.get("role", ""), "accountType": item.get("accountType", ""), "customers": "All customers" if item.get("companyIds") == ["*"] else [company_names.get(value, value) for value in item.get("companyIds", [])], "groups": item.get("groupIds", [])})
    elif report_id == "integration-health":
        columns = _columns(("name", "Integration"), ("type", "Provider"), ("enabled", "Enabled"), ("status", "Connection status"), ("lastRun", "Last run"), ("runStatus", "Run status"), ("message", "Run detail"))
        latest = {}
        for run in sync_runs:
            latest.setdefault(run.get("type"), run)
        for item in integrations:
            run = latest.get(item.get("type"), {})
            rows.append({"name": item.get("name", ""), "type": item.get("type", ""), "enabled": item.get("enabled", False), "status": item.get("status", ""), "lastRun": run.get("finishedAt") or run.get("startedAt", ""), "runStatus": run.get("status", "Never run"), "message": run.get("message", "")})
    else:
        raise ValueError("Unknown report")

    return {
        "id": report_id,
        "title": definition["title"],
        "description": definition["description"],
        "generatedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "companyId": company_id,
        "columns": columns,
        "rows": rows,
        "summary": {"rowCount": len(rows), "customerCount": len({row.get("customer") for row in rows if row.get("customer")})},
    }


def report_filename(report: dict, extension: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", report["title"].casefold()).strip("-")
    return f"{slug}-{report['generatedAt'][:10]}.{extension}"


def render_csv(report: dict) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=[item["key"] for item in report["columns"]], extrasaction="ignore")
    writer.writerow({item["key"]: item["label"] for item in report["columns"]})
    for row in report["rows"]:
        writer.writerow({key: _text(value) for key, value in row.items()})
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def render_xlsx(report: dict, branding: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    accent = str(branding.get("accent", "#50d5b9")).lstrip("#")
    sheet.append([branding.get("name", "CMDB Hub"), report["title"]])
    sheet.append(["Generated UTC", report["generatedAt"]])
    sheet.append([])
    sheet.append([column["label"] for column in report["columns"]])
    for cell in sheet[4]:
        cell.fill = PatternFill("solid", fgColor=accent)
        cell.font = Font(bold=True, color="07131D")
    for row in report["rows"]:
        sheet.append([_text(row.get(column["key"], "")) for column in report["columns"]])
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(max(1, len(report['columns'])))}{max(4, sheet.max_row)}"
    for index, column in enumerate(report["columns"], 1):
        values = [_text(row.get(column["key"], "")) for row in report["rows"][:300]]
        sheet.column_dimensions[get_column_letter(index)].width = min(52, max(12, len(column["label"]) + 2, *(len(value) + 2 for value in values)))
    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_pdf(report: dict, branding: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm)
    styles = getSampleStyleSheet()
    accent = colors.HexColor(branding.get("accent", "#50d5b9"))
    story = [Paragraph(str(branding.get("name", "CMDB Hub")), styles["Heading3"]), Paragraph(report["title"], styles["Title"]), Paragraph(f"Generated {report['generatedAt']} · {report['summary']['rowCount']} rows", styles["Normal"]), Spacer(1, 6 * mm)]
    table_data = [[Paragraph(column["label"], styles["BodyText"]) for column in report["columns"]]]
    for row in report["rows"][:500]:
        table_data.append([Paragraph(_text(row.get(column["key"], ""))[:500].replace("&", "&amp;").replace("<", "&lt;"), styles["BodyText"]) for column in report["columns"]])
    widths = [(landscape(A4)[0] - 24 * mm) / max(1, len(report["columns"]))] * len(report["columns"])
    table = Table(table_data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), accent), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#07131D")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BCC6D4")), ("FONTSIZE", (0, 0), (-1, -1), 7), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")])]))
    story.append(table)
    document.build(story)
    return output.getvalue()
